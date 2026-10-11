from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import os
import sqlite3
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from xml.sax.saxutils import escape as xml_escape


MAX_JSON_BYTES = 25 * 1024 * 1024
MAX_FINDINGS = 5000
MAX_SOURCES_PER_FINDING = 100


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds").replace(
        "+00:00", "Z"
    )


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while True:
            chunk = stream.read(1024 * 1024)
            if not chunk:
                break
            digest.update(chunk)
    return digest.hexdigest()


def load_json(path_text: str) -> tuple[dict | list, dict]:
    path = Path(path_text).expanduser().resolve(strict=True)
    if not path.is_file():
        raise ValueError(f"Not a file: {path}")
    if path.stat().st_size > MAX_JSON_BYTES:
        raise ValueError(f"JSON input exceeds {MAX_JSON_BYTES} bytes: {path}")

    raw = path.read_bytes()
    try:
        data = json.loads(raw.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValueError(f"Invalid UTF-8 JSON: {path}") from exc

    return data, {
        "path": str(path),
        "sha256": hashlib.sha256(raw).hexdigest(),
        "bytes": len(raw),
    }


def read_case_database(db_path_text: str, case_id: str) -> tuple[list, list]:
    db_path = Path(db_path_text).expanduser().resolve(strict=True)
    if not db_path.is_file():
        raise ValueError(f"Database is not a file: {db_path}")

    uri = db_path.as_uri() + "?mode=ro"
    conn = sqlite3.connect(uri, uri=True)
    conn.row_factory = sqlite3.Row

    try:
        evidence = conn.execute(
            """
            SELECT id, case_id, original_name, extension, notes,
                   byte_size, sha256, sha3_256, received_at, received_by
            FROM evidence
            WHERE case_id = ?
            ORDER BY received_at, id
            """,
            (case_id,),
        ).fetchall()

        evidence_ids = [row["id"] for row in evidence]
        custody = []

        if evidence_ids:
            placeholders = ",".join("?" for _ in evidence_ids)
            custody = conn.execute(
                f"""
                SELECT evidence_id, created_at, actor, action, custodian, purpose
                FROM custody
                WHERE evidence_id IN ({placeholders})
                ORDER BY evidence_id, id
                """,
                evidence_ids,
            ).fetchall()

        return [dict(row) for row in evidence], [dict(row) for row in custody]
    finally:
        conn.close()


def validate_findings(
    findings_data: dict | list | None,
    case_id: str,
    known_evidence_ids: set[str],
) -> list[dict]:
    if findings_data is None:
        return []

    if isinstance(findings_data, dict):
        declared_case = findings_data.get("case_id")
        if declared_case and declared_case != case_id:
            raise ValueError("Findings JSON case_id does not match requested case.")
        findings = findings_data.get("findings", [])
    elif isinstance(findings_data, list):
        findings = findings_data
    else:
        raise ValueError("Findings input must be a JSON object or array.")

    if not isinstance(findings, list) or len(findings) > MAX_FINDINGS:
        raise ValueError(f"Findings must be a list with at most {MAX_FINDINGS} items.")

    allowed_classes = {"observation", "interpretation", "lead"}
    allowed_review = {"reviewed", "unreviewed"}
    normalized = []

    for index, raw in enumerate(findings, start=1):
        if not isinstance(raw, dict):
            raise ValueError(f"Finding {index} must be a JSON object.")

        finding = dict(raw)
        for required in ("id", "title", "description"):
            if not isinstance(finding.get(required), str) or not finding[required].strip():
                raise ValueError(f"Finding {index} requires a non-empty {required}.")

        finding.setdefault("classification", "lead")
        finding.setdefault("review_status", "unreviewed")
        finding.setdefault("reviewer", "")
        finding.setdefault("benign_explanations", [])
        finding.setdefault("sources", [])

        if finding["classification"] not in allowed_classes:
            raise ValueError(
                f"Finding {index} classification must be one of "
                f"{sorted(allowed_classes)}."
            )
        if finding["review_status"] not in allowed_review:
            raise ValueError(
                f"Finding {index} review_status must be reviewed or unreviewed."
            )
        if finding["review_status"] == "reviewed" and not finding["reviewer"].strip():
            raise ValueError(f"Reviewed finding {index} needs a reviewer name.")

        if not isinstance(finding["sources"], list):
            raise ValueError(f"Finding {index} sources must be a list.")
        if len(finding["sources"]) > MAX_SOURCES_PER_FINDING:
            raise ValueError(f"Finding {index} has too many sources.")

        clean_sources = []
        for source_index, source in enumerate(finding["sources"], start=1):
            if not isinstance(source, dict):
                raise ValueError(
                    f"Finding {index}, source {source_index} must be an object."
                )
            evidence_id = source.get("evidence_id")
            locator = source.get("source_locator")
            if evidence_id not in known_evidence_ids:
                raise ValueError(
                    f"Finding {index} references unknown evidence ID {evidence_id!r}."
                )
            if not isinstance(locator, str) or not locator.strip():
                raise ValueError(
                    f"Finding {index}, source {source_index} needs source_locator."
                )
            clean_sources.append(
                {
                    "evidence_id": evidence_id,
                    "source_locator": locator[:1000],
                    "anchor_text": str(source.get("anchor_text", ""))[:4000],
                    "exporter_verification": "not verified by report exporter",
                }
            )

        benign = finding["benign_explanations"]
        if not isinstance(benign, list):
            raise ValueError(f"Finding {index} benign_explanations must be a list.")

        finding["sources"] = clean_sources
        finding["benign_explanations"] = [
            str(item)[:2000] for item in benign[:100]
        ]
        normalized.append(finding)

    return normalized


def summarize_run(kind: str, path_text: str) -> dict:
    payload, file_info = load_json(path_text)
    if not isinstance(payload, dict):
        raise ValueError(f"{kind} manifest must be a JSON object: {path_text}")

    summary = {
        "kind": kind,
        "manifest_path": file_info["path"],
        "manifest_sha256": file_info["sha256"],
        "manifest_bytes": file_info["bytes"],
        "status": payload.get("status", "not stated"),
        "tool": payload.get("parser") or payload.get("volatility_executable") or kind,
        "created_at": payload.get("created_at_utc") or payload.get("created_at"),
        "input_sha256": (
            payload.get("input_sha256")
            or payload.get("source_image_hash_before", {}).get("sha256")
            or payload.get("memory_image_hash_before", {}).get("sha256")
        ),
        "warnings": payload.get("warnings", []),
    }

    if kind == "parser":
        summary["records_reported"] = len(payload.get("records", []))
        summary["parser_version"] = payload.get("parser_version")
    elif kind == "memory":
        outputs = payload.get("outputs", [])
        summary["plugin_statuses"] = [
            {
                "plugin": item.get("plugin"),
                "status": item.get("status"),
                "return_code": item.get("return_code"),
            }
            for item in outputs
        ]
    elif kind == "disk":
        summary["recovered_file_count"] = payload.get("recovered_file_count")
        summary["source_image_changed"] = payload.get("source_image_changed")

    return summary


def write_csv(path: Path, columns: list[str], rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            cooked = {}
            for column in columns:
                value = row.get(column, "")
                if isinstance(value, (dict, list)):
                    value = json.dumps(value, ensure_ascii=False, sort_keys=True)
                cooked[column] = value
            writer.writerow(cooked)


def build_html(report: dict) -> str:
    def esc(value) -> str:
        return html.escape(str(value or ""), quote=True)

    evidence_rows = "".join(
        "<tr>"
        f"<td>{esc(item['id'])}</td>"
        f"<td>{esc(item['original_name'])}</td>"
        f"<td>{esc(item['byte_size'])}</td>"
        f"<td><code>{esc(item['sha256'])}</code></td>"
        f"<td>{esc(item['received_at'])}</td>"
        "</tr>"
        for item in report["evidence"]
    )

    finding_sections = []
    for finding in report["findings"]:
        sources = "".join(
            "<li>"
            f"Exhibit {esc(src['evidence_id'])}; "
            f"{esc(src['source_locator'])}; "
            "anchor text not verified by exporter"
            "</li>"
            for src in finding["sources"]
        )
        benign = "".join(
            f"<li>{esc(item)}</li>" for item in finding["benign_explanations"]
        )
        finding_sections.append(
            "<section>"
            f"<h3>{esc(finding['id'])}: {esc(finding['title'])}</h3>"
            f"<p><b>Class:</b> {esc(finding['classification'])}; "
            f"<b>Review:</b> {esc(finding['review_status'])}; "
            f"<b>Reviewer:</b> {esc(finding['reviewer'])}</p>"
            f"<p>{esc(finding['description'])}</p>"
            f"<p><b>Sources</b></p><ul>{sources}</ul>"
            f"<p><b>Benign explanations recorded</b></p><ul>{benign}</ul>"
            "</section>"
        )

    return f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{esc(report['title'])}</title>
<style>
body{{font:15px system-ui,sans-serif;max-width:1100px;margin:2rem auto;padding:0 1rem;color:#191919}}
table{{border-collapse:collapse;width:100%;font-size:12px}}
th,td{{border:1px solid #bbb;padding:.45rem;text-align:left;vertical-align:top}}
code{{overflow-wrap:anywhere}} section{{border-top:1px solid #bbb;padding:1rem 0}}
.warning{{border:2px solid #a33;padding:.75rem;background:#fff8ee}}
@media print{{body{{max-width:none;margin:0}}}}
</style>
</head>
<body>
<h1>{esc(report['title'])}</h1>
<p><b>Case:</b> {esc(report['case_id'])}<br>
<b>Generated UTC:</b> {esc(report['generated_at_utc'])}<br>
<b>Examiner:</b> {esc(report['examiner'])}</p>
<div class="warning">{esc(report['disclaimer'])}</div>
<h2>Evidence inventory</h2>
<table><thead><tr><th>ID</th><th>Filename</th><th>Bytes</th>
<th>SHA-256</th><th>Received</th></tr></thead>
<tbody>{evidence_rows}</tbody></table>
<h2>Custody events</h2>
<pre>{esc(json.dumps(report['custody'], ensure_ascii=False, indent=2))}</pre>
<h2>Analysis runs</h2>
<pre>{esc(json.dumps(report['analysis_runs'], ensure_ascii=False, indent=2))}</pre>
<h2>Findings recorded by examiner</h2>
{''.join(finding_sections) if finding_sections else '<p>No findings supplied.</p>'}
<h2>Input file hashes</h2>
<pre>{esc(json.dumps(report['input_files'], ensure_ascii=False, indent=2))}</pre>
</body></html>"""


def build_docx(path: Path, report: dict) -> None:
    doc = Document()
    heading = doc.add_heading(report["title"], 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"Case: {report['case_id']}")
    doc.add_paragraph(f"Examiner: {report['examiner']}")
    doc.add_paragraph(f"Generated UTC: {report['generated_at_utc']}")
    doc.add_heading("Important limitation", level=1)
    doc.add_paragraph(report["disclaimer"])

    doc.add_heading("Evidence inventory", level=1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["Evidence ID", "Filename", "Bytes", "SHA-256", "Received UTC"]
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    for item in report["evidence"]:
        cells = table.add_row().cells
        values = [
            item["id"],
            item["original_name"],
            str(item["byte_size"]),
            item["sha256"],
            item["received_at"],
        ]
        for cell, text in zip(cells, values):
            cell.text = str(text)

    doc.add_heading("Custody events", level=1)
    if report["custody"]:
        table = doc.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        headers = ["Evidence ID", "Time UTC", "Actor", "Action", "Custodian", "Purpose"]
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
        for item in report["custody"]:
            cells = table.add_row().cells
            values = [
                item["evidence_id"], item["created_at"], item["actor"],
                item["action"], item["custodian"], item["purpose"],
            ]
            for cell, text in zip(cells, values):
                cell.text = str(text)
    else:
        doc.add_paragraph("No custody events found for this case.")

    doc.add_heading("Analysis runs", level=1)
    if report["analysis_runs"]:
        for run in report["analysis_runs"]:
            doc.add_heading(f"{run['kind']}: {run['tool']}", level=2)
            doc.add_paragraph(
                f"Status: {run['status']}; manifest SHA-256: "
                f"{run['manifest_sha256']}; source/input SHA-256: "
                f"{run.get('input_sha256') or 'not supplied'}"
            )
            doc.add_paragraph(f"Manifest path: {run['manifest_path']}")
            doc.add_paragraph(
                "Run summary: " + json.dumps(run, ensure_ascii=False)
            )
    else:
        doc.add_paragraph("No analysis manifests supplied.")

    doc.add_heading("Examiner-entered findings", level=1)
    if report["findings"]:
        for finding in report["findings"]:
            doc.add_heading(
                f"{finding['id']}: {finding['title']} "
                f"[{finding['review_status'].upper()}]",
                level=2,
            )
            doc.add_paragraph(
                f"Classification: {finding['classification']}; "
                f"reviewer: {finding['reviewer'] or 'not stated'}"
            )
            doc.add_paragraph(finding["description"])
            doc.add_paragraph("Sources:")
            if finding["sources"]:
                for source in finding["sources"]:
                    doc.add_paragraph(
                        f"Exhibit {source['evidence_id']} — "
                        f"{source['source_locator']}; "
                        "anchor text is not verified by exporter."
                    )
            else:
                doc.add_paragraph("No sources were recorded.")
            doc.add_paragraph("Benign explanations recorded:")
            for item in finding["benign_explanations"]:
                doc.add_paragraph(item, style="List Bullet")
    else:
        doc.add_paragraph("No findings supplied.")

    doc.add_heading("Input hashes", level=1)
    for item in report["input_files"]:
        doc.add_paragraph(
            f"{item['path']} — SHA-256 {item['sha256']} ({item['bytes']} bytes)"
        )

    doc.save(path)


def build_pdf(path: Path, report: dict) -> None:
    styles = getSampleStyleSheet()
    story = [
        Paragraph(xml_escape(report["title"]), styles["Title"]),
        Spacer(1, 0.15 * inch),
        Paragraph(f"Case: {xml_escape(report['case_id'])}", styles["Normal"]),
        Paragraph(f"Examiner: {xml_escape(report['examiner'])}", styles["Normal"]),
        Paragraph(
            f"Generated UTC: {xml_escape(report['generated_at_utc'])}",
            styles["Normal"],
        ),
        Spacer(1, 0.15 * inch),
        Paragraph(xml_escape(report["disclaimer"]), styles["BodyText"]),
        Spacer(1, 0.2 * inch),
        Paragraph("Evidence inventory", styles["Heading1"]),
    ]

    evidence_data = [["ID", "Filename", "Bytes", "SHA-256", "Received UTC"]]
    for item in report["evidence"]:
        evidence_data.append(
            [
                Paragraph(xml_escape(item["id"]), styles["BodyText"]),
                Paragraph(xml_escape(item["original_name"]), styles["BodyText"]),
                str(item["byte_size"]),
                Paragraph(xml_escape(item["sha256"]), styles["BodyText"]),
                Paragraph(xml_escape(item["received_at"]), styles["BodyText"]),
            ]
        )

    evidence_table = Table(
        evidence_data,
        colWidths=[0.85 * inch, 1.55 * inch, 0.55 * inch, 3.0 * inch, 1.3 * inch],
        repeatRows=1,
    )
    evidence_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e9e6df")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(evidence_table)

    story.extend(
        [
            Spacer(1, 0.2 * inch),
            Paragraph("Custody events", styles["Heading1"]),
        ]
    )
    if report["custody"]:
        for item in report["custody"]:
            text = (
                f"{item['created_at']} — {item['action']} — "
                f"{item['actor']} — custodian {item['custodian']} — "
                f"{item['purpose']}"
            )
            story.append(Paragraph(xml_escape(text), styles["BodyText"]))
    else:
        story.append(Paragraph("No custody events found.", styles["BodyText"]))

    story.append(Paragraph("Analysis runs", styles["Heading1"]))
    for run in report["analysis_runs"]:
        story.append(
            Paragraph(
                xml_escape(
                    f"{run['kind']}: {run['tool']}; status={run['status']}; "
                    f"manifest SHA-256={run['manifest_sha256']}; "
                    f"input SHA-256={run.get('input_sha256') or 'not supplied'}"
                ),
                styles["BodyText"],
            )
        )
        story.append(
            Paragraph(
                xml_escape(f"Manifest: {run['manifest_path']}"),
                styles["BodyText"],
            )
        )
    if not report["analysis_runs"]:
        story.append(Paragraph("No analysis manifests supplied.", styles["BodyText"]))

    story.append(Paragraph("Examiner-entered findings", styles["Heading1"]))
    for finding in report["findings"]:
        block = [
            Paragraph(
                xml_escape(
                    f"{finding['id']}: {finding['title']} "
                    f"[{finding['review_status'].upper()}]"
                ),
                styles["Heading2"],
            ),
            Paragraph(
                xml_escape(
                    f"Classification: {finding['classification']}; "
                    f"reviewer: {finding['reviewer'] or 'not stated'}"
                ),
                styles["BodyText"],
            ),
            Paragraph(xml_escape(finding["description"]), styles["BodyText"]),
        ]
        for source in finding["sources"]:
            block.append(
                Paragraph(
                    xml_escape(
                        f"Source: exhibit {source['evidence_id']}, "
                        f"{source['source_locator']}; anchor not verified by exporter."
                    ),
                    styles["BodyText"],
                )
            )
        story.append(KeepTogether(block))

    if not report["findings"]:
        story.append(Paragraph("No findings supplied.", styles["BodyText"]))

    story.append(Paragraph("Input file hashes", styles["Heading1"]))
    for item in report["input_files"]:
        story.append(
            Paragraph(
                xml_escape(
                    f"{item['path']} — SHA-256 {item['sha256']} "
                    f"({item['bytes']} bytes)"
                ),
                styles["BodyText"],
            )
        )

    doc = SimpleDocTemplate(
        str(path),
        pagesize=landscape(letter),
        rightMargin=0.45 * inch,
        leftMargin=0.45 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
        title=report["title"],
        author=report["examiner"],
    )
    doc.build(story)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Create transparent case reports from Phase 1 and analysis outputs."
    )
    parser.add_argument("--db", required=True, help="Phase 1 SQLite database.")
    parser.add_argument("--case-id", required=True)
    parser.add_argument("--examiner", required=True)
    parser.add_argument("--title", default="Digital Forensic Case Report")
    parser.add_argument("--findings-json")
    parser.add_argument("--parser-result", action="append", default=[])
    parser.add_argument("--disk-manifest", action="append", default=[])
    parser.add_argument("--memory-manifest", action="append", default=[])
    parser.add_argument("--integrity-result")
    parser.add_argument("--output-root", required=True)
    args = parser.parse_args()

    evidence, custody = read_case_database(args.db, args.case_id)
    if not evidence:
        raise ValueError(f"No evidence found for case_id={args.case_id!r}.")

    input_files = []
    findings_data = None

    if args.findings_json:
        findings_data, info = load_json(args.findings_json)
        input_files.append(info)

    analysis_runs = []
    for kind, paths in (
        ("parser", args.parser_result),
        ("disk", args.disk_manifest),
        ("memory", args.memory_manifest),
    ):
        for path in paths:
            summary = summarize_run(kind, path)
            analysis_runs.append(summary)
            input_files.append(
                {
                    "path": summary["manifest_path"],
                    "sha256": summary["manifest_sha256"],
                    "bytes": summary["manifest_bytes"],
                }
            )

    integrity_result = None
    if args.integrity_result:
        integrity_result, info = load_json(args.integrity_result)
        input_files.append(info)

    evidence_ids = {item["id"] for item in evidence}
    findings = validate_findings(findings_data, args.case_id, evidence_ids)

    report = {
        "schema": "chronofact-phase5-report-v1",
        "title": args.title,
        "case_id": args.case_id,
        "examiner": args.examiner,
        "generated_at_utc": utc_now(),
        "disclaimer": (
            "This is an exporter-generated report, not a statutory certificate, "
            "legal opinion, or automated determination of credibility. Examiner "
            "findings and anchor text are user-supplied. The exporter checks that "
            "referenced evidence IDs exist but does not verify quotations against "
            "the original evidence bytes."
        ),
        "evidence": evidence,
        "custody": custody,
        "analysis_runs": analysis_runs,
        "integrity_result": integrity_result,
        "findings": findings,
        "input_files": input_files,
    }

    root = Path(args.output_root).expanduser().resolve()
    root.mkdir(parents=True, exist_ok=True)
    output_dir = root / (
        f"report-{datetime.now().strftime('%Y%m%d-%H%M%S')}-{uuid.uuid4().hex[:8]}"
    )
    output_dir.mkdir(parents=False, exist_ok=False)

    (output_dir / "case_report.json").write_text(
        json.dumps(report, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )

    write_csv(
        output_dir / "evidence.csv",
        [
            "id", "case_id", "original_name", "extension", "notes",
            "byte_size", "sha256", "sha3_256", "received_at", "received_by",
        ],
        evidence,
    )
    write_csv(
        output_dir / "custody.csv",
        [
            "evidence_id", "created_at", "actor", "action",
            "custodian", "purpose",
        ],
        custody,
    )

    flat_findings = []
    flat_sources = []
    for finding in findings:
        flat_findings.append(
            {
                "id": finding["id"],
                "title": finding["title"],
                "classification": finding["classification"],
                "review_status": finding["review_status"],
                "reviewer": finding["reviewer"],
                "description": finding["description"],
                "benign_explanations": finding["benign_explanations"],
            }
        )
        for source in finding["sources"]:
            flat_sources.append({"finding_id": finding["id"], **source})

    write_csv(
        output_dir / "findings.csv",
        [
            "id", "title", "classification", "review_status",
            "reviewer", "description", "benign_explanations",
        ],
        flat_findings,
    )
    write_csv(
        output_dir / "finding_sources.csv",
        [
            "finding_id", "evidence_id", "source_locator",
            "anchor_text", "exporter_verification",
        ],
        flat_sources,
    )
    write_csv(
        output_dir / "analysis_runs.csv",
        [
            "kind", "tool", "status", "created_at", "input_sha256",
            "manifest_path", "manifest_sha256",
        ],
        analysis_runs,
    )

    (output_dir / "case_report.html").write_text(
        build_html(report), encoding="utf-8"
    )
    build_docx(output_dir / "case_report.docx", report)
    build_pdf(output_dir / "case_report.pdf", report)

    outputs = []
    for path in sorted(output_dir.iterdir()):
        if path.is_file() and path.name != "output_manifest.json":
            outputs.append(
                {
                    "file": path.name,
                    "bytes": path.stat().st_size,
                    "sha256": sha256_file(path),
                }
            )

    (output_dir / "output_manifest.json").write_text(
        json.dumps(
            {
                "schema": "chronofact-phase5-output-manifest-v1",
                "created_at_utc": utc_now(),
                "files": outputs,
                "note": (
                    "Local output checksums only; not a digital signature "
                    "or trusted timestamp."
                ),
            },
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )

    print(f"Report package created: {output_dir}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"Report generation failed: {type(exc).__name__}: {exc}", file=sys.stderr)
        raise SystemExit(1)